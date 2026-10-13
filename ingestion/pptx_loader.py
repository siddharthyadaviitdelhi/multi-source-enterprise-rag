"""PowerPoint (.pptx) ingestion: extract text from each slide's text boxes and tables."""
from pathlib import Path
from pptx import Presentation
from .common import Document, clean_text


def load_pptx(path) -> list[Document]:
    """Return one Document per slide that contains extractable text."""
    path = Path(path)
    presentation = Presentation(str(path))
    docs = []

    for slide_num, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            # Extract ordinary text boxes and titles.
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text
                if text and text.strip():
                    parts.append(text.strip())

            # Extract table cells, if the slide contains a table.
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))

        text = clean_text("\\n".join(parts))
        if text:
            docs.append(
                Document(
                    text=text,
                    source=path.name,
                    source_type="pptx",
                    metadata={"slide": slide_num, "path": str(path)},
                )
            )

    return docs
